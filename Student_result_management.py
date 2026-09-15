import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE = "student_results.xlsx"

# Create Excel file
if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Roll No.", "Class", "Sub1", "Sub2", "Sub3",
               "Sub4", "Sub5", "Total", "Percentage", "Result"])
    wb.save(FILE)


def save_student():
    try:
        name = name_entry.get()
        roll = int(roll_entry.get())
        cls = class_entry.get()

        marks = [int(e.get()) for e in subject_entries]

        if any(m < 0 or m > 100 for m in marks):
            messagebox.showerror("Error", "Marks must be 0 to 100")
            return

        total = sum(marks)
        percentage = total / 5
        result = "Pass" if all(m >= 40 for m in marks) else "Fail"

        wb = load_workbook(FILE)
        ws = wb.active
        ws.append([name, roll, cls, *marks, total, percentage, result])
        wb.save(FILE)

        messagebox.showinfo(
            "Success",
            f"Record Saved!\nTotal: {total}\nPercentage: {percentage:.2f}%\nResult: {result}"
        )

    except ValueError:
        messagebox.showerror("Error", "Please enter valid details")


def get_result():
    roll = result_entry.get()

    for item in result_tree.get_children():
        result_tree.delete(item)

    wb = load_workbook(FILE, data_only=True)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == roll:
            result_tree.insert("", "end",
                values=(row[0], row[1], row[2], row[8],
                        f"{row[9]:.2f}%", row[10]))
            return

    messagebox.showerror("Error", "Student record not found")


def show_all():
    for item in all_tree.get_children():
        all_tree.delete(item)

    wb = load_workbook(FILE, data_only=True)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        all_tree.insert("", "end",
            values=(row[0], row[1], row[2], row[8],
                    f"{row[9]:.2f}%", row[10]))


# Main Window
root = tk.Tk()
root.title("Student Result Management System")
root.geometry("800x600")

tk.Label(root, text="STUDENT RESULT MANAGEMENT SYSTEM",
         font=("Arial", 18, "bold")).pack(pady=10)

# Add Student
frame = tk.LabelFrame(root, text="Add Student", padx=10, pady=10)
frame.pack(padx=15, pady=5, fill="x")

tk.Label(frame, text="Name").grid(row=0, column=0)
name_entry = tk.Entry(frame)
name_entry.grid(row=0, column=1)

tk.Label(frame, text="Roll No.").grid(row=0, column=2)
roll_entry = tk.Entry(frame)
roll_entry.grid(row=0, column=3)

tk.Label(frame, text="Class").grid(row=0, column=4)
class_entry = tk.Entry(frame)
class_entry.grid(row=0, column=5)

subject_entries = []

for i in range(5):
    tk.Label(frame, text=f"Sub {i+1}").grid(row=1, column=i)
    e = tk.Entry(frame, width=10)
    e.grid(row=2, column=i, padx=5)
    subject_entries.append(e)

tk.Button(frame, text="💾 Save", command=save_student,
          width=15).grid(row=3, column=0, columnspan=3, pady=10)


# Get Result
get_frame = tk.LabelFrame(root, text="Get Result", padx=10, pady=10)
get_frame.pack(padx=15, pady=5, fill="x")

tk.Label(get_frame, text="Enter Roll No.").pack(side="left")
result_entry = tk.Entry(get_frame)
result_entry.pack(side="left", padx=10)

tk.Button(get_frame, text="🔍 Get Result",
          command=get_result).pack(side="left")


columns = ("Name", "Roll No.", "Class", "Total", "Percentage", "Result")

result_tree = ttk.Treeview(get_frame, columns=columns,
                           show="headings", height=2)

for col in columns:
    result_tree.heading(col, text=col)
    result_tree.column(col, width=110)

result_tree.pack(pady=10)


# Show All Results
tk.Button(root, text="📋 Show All Results",
          command=show_all, width=20).pack(pady=10)

all_tree = ttk.Treeview(root, columns=columns,
                        show="headings", height=8)

for col in columns:
    all_tree.heading(col, text=col)
    all_tree.column(col, width=120)

all_tree.pack(padx=15, pady=5)

root.mainloop()
